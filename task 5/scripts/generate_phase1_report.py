"""Generate Phase 1 Review Report for CortexHR AI/ML in both PDF and DOCX formats.
Follows all prompt guidelines: exactly 2-3 pages, professional layout, honest evidence,
natural student voice, real numbers and file paths.
"""
import os
import sys
from pathlib import Path

# ---------- REPORT CONTENT & DATA ----------
REPORT_HEADER = {
    "title": "Phase 1 Review Report — CortexHR AI/ML Employee Attrition",
    "name": "Dikshit Kumar",
    "track": "AI/ML",
    "phase": "Phase 1",
    "project": "CortexHR AI/ML — Employee Attrition",
    "repo": "https://github.com/dikshit0012/CortexHR-.git",
    "branch": "master",
    "tag": "phase-1",
}

TASKS_SUMMARY = [
    {
        "task": "Task 1: Data Audit & Exploration",
        "status": "Complete & Tested",
        "worked_on": "Explored raw dataset (2,389 rows, 42 columns). Detected 6 duplicate rows, 4 negative tenure errors, and 6 severe target leakage columns (resignation date, exit reason, etc.). Assessed class imbalance (14.7% attrition rate).",
        "evidence": "Task 1/cortexhr-ml/reports/findings.json; Task1_Data_Audit_Report.pdf",
    },
    {
        "task": "Task 2: Reusable Feature Pipeline",
        "status": "Complete & Tested",
        "worked_on": "Engineered clean scikit-learn ColumnTransformer pipeline. Dropped duplicates and invalid tenure rows to yield 2,379 clean employee records. Joined 3-month rolling attendance/payroll aggregates strictly inside cross-validation folds.",
        "evidence": "task 2/files/extracted/cortexhr-ml-task5/reports/tables/day2_pipeline_run.json; src/cortexhr/features.py",
    },
    {
        "task": "Task 3: Train Honest Baselines",
        "status": "Complete & Tested",
        "worked_on": "Trained baseline models using 5-fold stratified CV: Dummy (always stays), balanced Logistic Regression, and Decision Trees. Validated accuracy paradox under class imbalance and diagnosed severe decision tree overfitting.",
        "evidence": "task 4/files/cortexhr-ml/reports/tables/day3_baseline_comparison.csv; reports/day3_baselines.md; ROC-AUC: 0.712 vs 0.713 benchmark (2,379 real rows — not the 2,400-row synthetic stand-in)",
    },
    {
        "task": "Task 4: Evaluation & Risk Bands",
        "status": "Complete & Tested",
        "worked_on": "Evaluated out-of-fold predictions. Calibrated probabilities with sigmoid scaling (Brier score 0.206 -> 0.116). Calculated top-decile lift (2.26x) and defined 3 actionable operational risk bands (High, Medium, Low). Authored executive summary.",
        "evidence": "task 4/files/day4_metrics_table.csv; task 4/files/Day4_Manager_Note.pdf; 29 unit tests passing",
    },
    {
        "task": "Task 5: Unsupervised Attendance Anomaly Detector",
        "status": "Complete & Tested",
        "worked_on": "Trained Isolation Forest (3% contamination) on 7,339 employee-months without labels. Combined with domain rules (unexcused absence, impossible hours). Measured against hidden labels using catch_rate(). Built HR priority tiers.",
        "evidence": "task 5/reports/task5_metrics.csv; task 5/reports/task5_missed_anomalies.csv; 14/14 unit tests passing",
    },
]

METRICS_SUMMARY = {
    "task1": {
        "raw_rows": "2,389",
        "clean_rows": "2,379",
        "duplicates": "6 dropped (EMP-0060, EMP-0237, EMP-0238, EMP-0325, EMP-0483, EMP-0570)",
        "invalid_tenure": "4 dropped (EMP-0248, EMP-0280, EMP-0568, EMP-0596)",
        "leaks": "6 columns excluded (resignation_date, exit_date, exit_reason, etc.)",
        "prevalence": "14.67% (349 leavers out of 2,379)",
    },
    "task3": [
        ("Dummy (always stays)", "0.500 +/- 0.000", "0.147", "0.000", "0.000", "0.853", "0.854 benchmark"),
        ("Logistic Regression (balanced)", "0.712 +/- 0.018", "0.309", "0.265", "0.633", "0.686", "0.713 benchmark"),
        ("Decision Tree (depth 4, balanced)", "0.654 +/- 0.038", "0.255", "0.236", "0.525", "0.681", "0.650 benchmark"),
        ("Decision Tree (unconstrained)", "0.542 +/- 0.028", "0.164", "0.214", "0.226", "0.765", "Overfit (train ROC: 0.99+)"),
    ],
    "task4": {
        "brier_uncalibrated": "0.206",
        "brier_calibrated": "0.116",
        "brier_naive_prevalence": "0.125",
        "top_decile_lift": "2.26x (33.2% leaver concentration vs 14.7% base prevalence)",
        "top_decile_recall": "22.6% of all leavers captured in the top 10% cutoff",
        "risk_bands": [
            ("High Risk", "Top 10% (238 employees)", "33.19%", "2.26x", "22.6%"),
            ("Medium Risk", "Next 20% (476 employees)", "23.74%", "1.62x", "32.4%"),
            ("Low Risk", "Bottom 70% (1,665 employees)", "9.43%", "0.64x", "45.0%"),
        ],
    },
    "task5": [
        ("Published Baseline (anomaly_baselines)", "0.403", "0.593", "0.480", "89", "132", "61", "221 (3.01%)"),
        ("Baseline Recipe Re-run (raw 10 cols, 3%)", "0.398", "0.587", "0.474", "88", "133", "62", "221 (3.01%)"),
        ("Isolation Forest (curated features)", "0.412", "0.607", "0.491", "91", "130", "59", "221 (3.01%)"),
        ("Rule-Based Only (domain rules)", "0.918", "0.600", "0.726", "90", "8", "60", "98 (1.34%)"),
        ("Combined (IF OR Business Rules)", "0.488", "0.827", "0.614", "124", "130", "26", "254 (3.46%)"),
    ],
}

LEARNING_REFLECTION = (
    "During Phase 1, I learned that practical machine learning is mostly about data discipline rather than complex algorithms. "
    "I saw firsthand how target leakage columns like resignation dates can produce deceptively perfect results if left in the dataset, "
    "and why an 85% accuracy baseline is completely misleading when only 14.7% of employees leave. "
    "Calibrating probabilities taught me that raw classification scores cannot be handed to business managers without adjustment, "
    "and building the attendance detector showed me how unsupervised isolation trees and simple business rules can catch anomalies effectively without ever training on ground-truth labels."
)

PROBLEMS_AND_SOLUTIONS = [
    {
        "issue": "1. Subtle Target Leakage in Raw Employee Features",
        "description": (
            "During Task 1 exploration, several columns like 'resignation_date', 'exit_date', and 'exit_reason' had over 85% missing values. "
            "Initially, it looked like standard missing data. However, checking the target label revealed that these columns were populated ONLY for employees who left the company."
        ),
        "solution": (
            "Keeping these columns would allow a model to achieve 100% training accuracy through pure hindsight leakage. "
            "I classified all 6 exit-related columns as 'LEAK' in the data dictionary and strictly excluded them from the feature pipeline."
        ),
    },
    {
        "issue": "2. High Training Score vs Collapsing Cross-Validation in Decision Trees",
        "description": (
            "When running baseline models in Task 3, an unconstrained Decision Tree appeared to learn the training split almost perfectly (training ROC-AUC > 0.99). "
            "However, its 5-fold cross-validation ROC-AUC collapsed to 0.542, performing barely better than random coin-flipping."
        ),
        "solution": (
            "I ran a depth-tuning experiment across max_depth values 1 through 8. Constraining max_depth to 4 prevented the tree from memorizing small sample idiosyncrasies, "
            "yielding a stable cross-validation ROC-AUC of 0.654 +/- 0.038 and matching the benchmark."
        ),
    },
    {
        "issue": "3. Single-Feature Anomaly Invisibility in Multi-Dimensional Isolation Forest",
        "description": (
            "In Task 5, our unsupervised Isolation Forest caught almost all absence streak, sudden drop, and impossible hours anomalies, but missed 26 out of 30 LATE_SPIKE records (13.3% recall). "
            "Examining the missed records revealed that employees with late arrival spikes had completely normal attendance rates (0.98) and standard working hours (8.55 hrs/day)."
        ),
        "solution": (
            "Because 8 out of 9 features were completely normal, random isolation partitioning did not isolate these records early enough to beat the 3% cut-off. "
            "I documented this as an inherent limitation of tree-based multi-feature isolation for one-dimensional spikes, and designed a priority ranking system that combines rule triggers with model scores to flag such cases for HR review."
        ),
    },
]

PHASE_2_PLAN = [
    "1. Reusing Established Pipeline Components: Port the clean ColumnTransformer structure and rolling window join logic from Task 2 into the Phase 2 payroll workflows to ensure zero data leakage.",
    "2. Payroll Feature Engineering: Construct monthly payroll indicators (such as overtime-to-base salary ratios, retroactive adjustments, and sudden compensation rate shifts) that capture financial anomalies.",
    "3. Dual-Layer Anomaly Detection on Payroll: Apply Isolation Forest paired with strict statutory/labour rules (such as minimum wage violations, negative net pay, or excessive monthly overtime hours).",
    "4. Rigorous Out-of-Fold Evaluation: Keep validation labels strictly sealed in the vault until predictions are finalized, evaluating detector performance via precision, recall, and financial catch rate.",
    "5. Actionable Payroll Review Tiers: Implement a points-based prioritization queue so payroll accountants can audit High-risk payroll batches first, minimizing financial error risk before disbursement.",
]


# ==============================================================================
# REPORTLAB PDF GENERATION
# ==============================================================================
def generate_pdf(output_path: str):
    from reportlab.lib import colors
    from reportlab.lib.pagesizes import letter
    from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
    from reportlab.lib.units import inch
    from reportlab.pdfgen import canvas
    from reportlab.platypus import (
        SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle, KeepTogether, HRFlowable
    )

    class NumberedCanvas(canvas.Canvas):
        def __init__(self, *args, **kwargs):
            super().__init__(*args, **kwargs)
            self._saved_page_states = []

        def showPage(self):
            self._saved_page_states.append(dict(self.__dict__))
            self._startPage()

        def save(self):
            num_pages = len(self._saved_page_states)
            print(f"[PDF Build] Exact Total Pages: {num_pages}")
            for state in self._saved_page_states:
                self.__dict__.update(state)
                self.draw_page_decorations(num_pages)
                super().showPage()
            super().save()

        def draw_page_decorations(self, page_count):
            self.saveState()
            self.setFont("Helvetica", 8)
            self.setFillColor(colors.HexColor("#555555"))
            
            # Header on pages 2+
            if self._pageNumber > 1:
                self.drawString(40, letter[1] - 30, "CortexHR AI/ML — Phase 1 Review Report | Dikshit Kumar")
                self.setStrokeColor(colors.HexColor("#D0D5DD"))
                self.setLineWidth(0.5)
                self.line(40, letter[1] - 34, letter[0] - 40, letter[1] - 34)

            # Footer on all pages
            page_text = f"Page {self._pageNumber} of {page_count}"
            self.drawRightString(letter[0] - 40, 25, page_text)
            self.drawString(40, 25, "Confidential — CortexHR Internship Review Submission")
            self.setStrokeColor(colors.HexColor("#D0D5DD"))
            self.setLineWidth(0.5)
            self.line(40, 35, letter[0] - 40, 35)
            self.restoreState()

    doc = SimpleDocTemplate(
        output_path,
        pagesize=letter,
        leftMargin=38,
        rightMargin=38,
        topMargin=42,
        bottomMargin=42,
    )

    styles = getSampleStyleSheet()

    # Custom styles
    title_style = ParagraphStyle(
        "DocTitle",
        fontName="Helvetica-Bold",
        fontSize=17,
        leading=21,
        textColor=colors.HexColor("#1A2B4C"),
        spaceAfter=4,
    )
    subtitle_style = ParagraphStyle(
        "DocSubtitle",
        fontName="Helvetica-Bold",
        fontSize=10,
        leading=14,
        textColor=colors.HexColor("#2B6CB0"),
        spaceAfter=8,
    )
    h1_style = ParagraphStyle(
        "SectionH1",
        fontName="Helvetica-Bold",
        fontSize=11.5,
        leading=15,
        textColor=colors.HexColor("#1A2B4C"),
        spaceBefore=8,
        spaceAfter=4,
        keepWithNext=True,
    )
    h2_style = ParagraphStyle(
        "SectionH2",
        fontName="Helvetica-Bold",
        fontSize=9.5,
        leading=13,
        textColor=colors.HexColor("#2C3E50"),
        spaceBefore=5,
        spaceAfter=3,
        keepWithNext=True,
    )
    body_style = ParagraphStyle(
        "BodyTextCustom",
        fontName="Helvetica",
        fontSize=8.5,
        leading=11.5,
        textColor=colors.HexColor("#2D3748"),
        spaceAfter=4,
    )
    body_bold = ParagraphStyle(
        "BodyBoldCustom",
        fontName="Helvetica-Bold",
        fontSize=8.5,
        leading=11.5,
        textColor=colors.HexColor("#2D3748"),
    )
    table_cell = ParagraphStyle(
        "TableCell",
        fontName="Helvetica",
        fontSize=7.5,
        leading=9.5,
        textColor=colors.HexColor("#2D3748"),
    )
    table_cell_bold = ParagraphStyle(
        "TableCellBold",
        fontName="Helvetica-Bold",
        fontSize=7.5,
        leading=9.5,
        textColor=colors.HexColor("#1A2B4C"),
    )
    table_header = ParagraphStyle(
        "TableHeader",
        fontName="Helvetica-Bold",
        fontSize=7.5,
        leading=9.5,
        textColor=colors.white,
    )
    meta_label = ParagraphStyle(
        "MetaLabel",
        fontName="Helvetica-Bold",
        fontSize=8,
        leading=10.5,
        textColor=colors.HexColor("#1A2B4C"),
    )
    meta_val = ParagraphStyle(
        "MetaVal",
        fontName="Helvetica",
        fontSize=8,
        leading=10.5,
        textColor=colors.HexColor("#2D3748"),
    )

    story = []

    # Title
    story.append(Paragraph(REPORT_HEADER["title"], title_style))
    story.append(Paragraph("Undergraduate Technical Review | B.Tech Computer Science | CortexHR Project", subtitle_style))
    story.append(HRFlowable(width="100%", thickness=1.5, color=colors.HexColor("#1A2B4C"), spaceBefore=1, spaceAfter=6))

    # Header metadata box (compact 2-column table)
    meta_data = [
        [
            Paragraph("<b>Student Name:</b>", meta_label), Paragraph(REPORT_HEADER["name"], meta_val),
            Paragraph("<b>Track / Phase:</b>", meta_label), Paragraph(f"{REPORT_HEADER['track']} | {REPORT_HEADER['phase']}", meta_val),
        ],
        [
            Paragraph("<b>Project:</b>", meta_label), Paragraph(REPORT_HEADER["project"], meta_val),
            Paragraph("<b>Git Branch:</b>", meta_label), Paragraph(f"{REPORT_HEADER['branch']} (Tag: pending)", meta_val),
        ],
        [
            Paragraph("<b>Repository:</b>", meta_label),
            Paragraph(f"{REPORT_HEADER['repo']}", meta_val),
            Paragraph("<b>Evaluation Integrity:</b>", meta_label),
            Paragraph("Labels strictly sealed during training", meta_val),
        ],
    ]
    meta_table = Table(meta_data, colWidths=[80, 200, 95, 160])
    meta_table.setStyle(TableStyle([
        ('BACKGROUND', (0,0), (-1,-1), colors.HexColor("#F7FAFC")),
        ('BOX', (0,0), (-1,-1), 0.5, colors.HexColor("#E2E8F0")),
        ('VALIGN', (0,0), (-1,-1), 'MIDDLE'),
        ('TOPPADDING', (0,0), (-1,-1), 2.5),
        ('BOTTOMPADDING', (0,0), (-1,-1), 2.5),
        ('LEFTPADDING', (0,0), (-1,-1), 4),
        ('RIGHTPADDING', (0,0), (-1,-1), 4),
    ]))
    story.append(meta_table)
    story.append(Spacer(1, 6))

    # SECTION 2: Summary of the Five Tasks
    story.append(Paragraph("1. Summary of Phase 1 Tasks", h1_style))
    story.append(Paragraph(
        "Phase 1 encompassed five interconnected milestones establishing a leak-free data audit, an automated feature engineering pipeline, rigorous baseline classification models, calibrated risk banding, and unsupervised attendance anomaly detection.",
        body_style
    ))

    tasks_table_data = [
        [
            Paragraph("Task", table_header),
            Paragraph("Status", table_header),
            Paragraph("What I Worked On", table_header),
            Paragraph("Verified Evidence", table_header),
        ]
    ]
    for row in TASKS_SUMMARY:
        tasks_table_data.append([
            Paragraph(f"<b>{row['task']}</b>", table_cell_bold),
            Paragraph(f"<font color='#0D6832'><b>{row['status']}</b></font>", table_cell),
            Paragraph(row["worked_on"], table_cell),
            Paragraph(f"<code>{row['evidence']}</code>", table_cell),
        ])

    tasks_table = Table(tasks_table_data, colWidths=[95, 75, 205, 160])
    tasks_table.setStyle(TableStyle([
        ('BACKGROUND', (0,0), (-1,0), colors.HexColor("#1A2B4C")),
        ('ALIGN', (0,0), (-1,-1), 'LEFT'),
        ('VALIGN', (0,0), (-1,-1), 'TOP'),
        ('ROWBACKGROUNDS', (0,1), (-1,-1), [colors.white, colors.HexColor("#F8FAFC")]),
        ('GRID', (0,0), (-1,-1), 0.5, colors.HexColor("#CBD5E1")),
        ('TOPPADDING', (0,0), (-1,-1), 3),
        ('BOTTOMPADDING', (0,0), (-1,-1), 3),
        ('LEFTPADDING', (0,0), (-1,-1), 4),
        ('RIGHTPADDING', (0,0), (-1,-1), 4),
    ]))
    story.append(tasks_table)
    story.append(Spacer(1, 7))

    # SECTION 3: Proof That the Work Runs
    story.append(Paragraph("2. Proof That the Work Runs & Verified Metrics", h1_style))
    story.append(Paragraph(
        "All code, test suites, and evaluation notebooks were executed against the real CortexHR dataset. No benchmark metrics were copied as my own; every figure reported below reflects verified local execution.",
        body_style
    ))

    # Task 1 & 2 Execution Proof
    story.append(Paragraph("A. Data Audit & Feature Pipeline (Tasks 1 & 2)", h2_style))
    story.append(Paragraph(
        f"• <b>Dataset Intake:</b> Audited 2,389 raw rows across 42 columns. Dropped 6 duplicate employee records (<code>{METRICS_SUMMARY['task1']['duplicates']}</code>) and excluded 4 impossible negative tenure entries (<code>{METRICS_SUMMARY['task1']['invalid_tenure']}</code>).<br/>"
        f"• <b>Target Leakage Protection:</b> Identified and dropped 6 exit-related columns (<code>{METRICS_SUMMARY['task1']['leaks']}</code>).<br/>"
        f"• <b>Pipeline Clean Output:</b> Yielded 2,379 clean records (349 leavers, 14.67% base prevalence). Engineered rolling 3-month attendance aggregates joined strictly inside CV folds to prevent temporal data leakage.",
        body_style
    ))

    # Task 3 Baseline Table
    story.append(Paragraph("B. Honest Baseline Classification (Task 3 — 5-Fold Stratified CV)", h2_style))
    t3_data = [
        [
            Paragraph("Model", table_header),
            Paragraph("ROC-AUC (Ours)", table_header),
            Paragraph("PR-AUC", table_header),
            Paragraph("Precision", table_header),
            Paragraph("Recall", table_header),
            Paragraph("Accuracy", table_header),
            Paragraph("Reference Benchmark", table_header),
        ]
    ]
    for row in METRICS_SUMMARY["task3"]:
        t3_data.append([
            Paragraph(f"<b>{row[0]}</b>", table_cell_bold),
            Paragraph(row[1], table_cell),
            Paragraph(row[2], table_cell),
            Paragraph(row[3], table_cell),
            Paragraph(row[4], table_cell),
            Paragraph(row[5], table_cell),
            Paragraph(row[6], table_cell),
        ])
    t3_table = Table(t3_data, colWidths=[120, 75, 50, 50, 50, 50, 140])
    t3_table.setStyle(TableStyle([
        ('BACKGROUND', (0,0), (-1,0), colors.HexColor("#2B6CB0")),
        ('GRID', (0,0), (-1,-1), 0.5, colors.HexColor("#CBD5E1")),
        ('ROWBACKGROUNDS', (0,1), (-1,-1), [colors.white, colors.HexColor("#F8FAFC")]),
        ('VALIGN', (0,0), (-1,-1), 'MIDDLE'),
        ('TOPPADDING', (0,0), (-1,-1), 2.5),
        ('BOTTOMPADDING', (0,0), (-1,-1), 2.5),
        ('LEFTPADDING', (0,0), (-1,-1), 4),
        ('RIGHTPADDING', (0,0), (-1,-1), 4),
    ]))
    story.append(t3_table)
    story.append(Spacer(1, 4))

    # Task 4 Calibration & Risk Bands
    story.append(Paragraph("C. Probability Calibration & Operational Risk Bands (Task 4)", h2_style))
    story.append(Paragraph(
        f"• <b>Calibration Impact:</b> Raw balanced Logistic Regression outputs ranking scores rather than true probabilities. Applying sigmoid calibration dropped the Brier score from <b>{METRICS_SUMMARY['task4']['brier_uncalibrated']}</b> to <b>{METRICS_SUMMARY['task4']['brier_calibrated']}</b>, substantially outperforming the naive base rate baseline (0.125).<br/>"
        f"• <b>Decile Lift:</b> The top 10% scored segment (238 employees) attained a <b>2.26x lift</b> (33.2% leaver concentration vs 14.7% baseline) capturing 22.6% of all eventual leavers.<br/>"
        f"• <b>Unit Tests:</b> All 29 evaluation test cases in <code>test_evaluation.py</code> executed cleanly.",
        body_style
    ))

    # Task 5 Anomaly Table
    story.append(Paragraph("D. Unsupervised Attendance Anomaly Detection (Task 5)", h2_style))
    t5_data = [
        [
            Paragraph("Detector Approach", table_header),
            Paragraph("Precision", table_header),
            Paragraph("Recall (Catch Rate)", table_header),
            Paragraph("F1 Score", table_header),
            Paragraph("TP", table_header),
            Paragraph("FP", table_header),
            Paragraph("FN", table_header),
            Paragraph("Flagged Records", table_header),
        ]
    ]
    for row in METRICS_SUMMARY["task5"]:
        is_comb = "Combined" in row[0]
        bg = colors.HexColor("#EFF6FF") if is_comb else colors.white
        t5_data.append([
            Paragraph(f"<b>{row[0]}</b>", table_cell_bold),
            Paragraph(row[1], table_cell),
            Paragraph(row[2], table_cell),
            Paragraph(row[3], table_cell),
            Paragraph(row[4], table_cell),
            Paragraph(row[5], table_cell),
            Paragraph(row[6], table_cell),
            Paragraph(row[7], table_cell),
        ])
    t5_table = Table(t5_data, colWidths=[150, 48, 70, 45, 30, 30, 30, 132])
    t5_table.setStyle(TableStyle([
        ('BACKGROUND', (0,0), (-1,0), colors.HexColor("#1A2B4C")),
        ('GRID', (0,0), (-1,-1), 0.5, colors.HexColor("#CBD5E1")),
        ('ROWBACKGROUNDS', (0,1), (-1,-1), [colors.white, colors.HexColor("#F8FAFC")]),
        ('VALIGN', (0,0), (-1,-1), 'MIDDLE'),
        ('TOPPADDING', (0,0), (-1,-1), 2.5),
        ('BOTTOMPADDING', (0,0), (-1,-1), 2.5),
        ('LEFTPADDING', (0,0), (-1,-1), 4),
        ('RIGHTPADDING', (0,0), (-1,-1), 4),
    ]))
    story.append(t5_table)
    story.append(Spacer(1, 4))
    story.append(Paragraph(
        "• <b>Evaluation Protocol:</b> Predictions were generated and written to disk without labels. Ground-truth labels were only loaded afterward inside <code>evaluation.py</code>.<br/>"
        "• <b>Anomaly Types Caught:</b> 100% recall on <code>ABSENCE_STREAK</code> (40/40), <code>IMPOSSIBLE_HOURS</code> (35/35), <code>SUDDEN_DROP</code> (30/30), and <code>ZERO_ATTENDANCE</code> (15/15).<br/>"
        "• <b>Unit Test Verification:</b> 14 of 14 unit tests passed in <code>tests/test_attendance_anomaly_detector.py</code>.",
        body_style
    ))
    story.append(Spacer(1, 7))

    # SECTION 4: What I Learned
    story.append(Paragraph("3. What I Learned (Student Reflection)", h1_style))
    story.append(Paragraph(
        f"<i>\"{LEARNING_REFLECTION}\"</i>",
        ParagraphStyle(
            "QuoteStyle",
            fontName="Helvetica-Oblique",
            fontSize=8.5,
            leading=12,
            textColor=colors.HexColor("#1A2B4C"),
            backColor=colors.HexColor("#F0F4F8"),
            borderColor=colors.HexColor("#2B6CB0"),
            borderWidth=1,
            borderPadding=6,
            spaceAfter=6,
        )
    ))
    story.append(Spacer(1, 4))

    # SECTION 5: Problems I Faced and How I Solved Them
    story.append(Paragraph("4. Real Engineering Challenges & Solutions", h1_style))
    for prob in PROBLEMS_AND_SOLUTIONS:
        story.append(Paragraph(f"<b>{prob['issue']}</b>", h2_style))
        story.append(Paragraph(f"<b>Problem:</b> {prob['description']}", body_style))
        story.append(Paragraph(f"<b>Investigation & Resolution:</b> {prob['solution']}", body_style))
        story.append(Spacer(1, 3))

    # SECTION 6: My Plan for Phase 2
    story.append(Paragraph("5. Roadmap for Phase 2 (Payroll Anomaly Detection)", h1_style))
    story.append(Paragraph(
        "In Phase 2, detection moves from attendance to payroll, where data errors translate directly into financial discrepancy and employee dissatisfaction. Building upon Phase 1 foundation:",
        body_style
    ))
    for item in PHASE_2_PLAN:
        story.append(Paragraph(f"• {item}", body_style))

    # Build PDF
    doc.build(story, canvasmaker=NumberedCanvas)
    print(f"[OK] PDF report successfully generated: {output_path}")


# ==============================================================================
# DOCX GENERATION
# ==============================================================================
def generate_docx(output_path: str):
    import docx
    from docx.shared import Inches, Pt, RGBColor
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.enum.table import WD_TABLE_ALIGNMENT
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn

    doc = docx.Document()

    # Set margins
    for sec in doc.sections:
        sec.top_margin = Inches(0.5)
        sec.bottom_margin = Inches(0.5)
        sec.left_margin = Inches(0.55)
        sec.right_margin = Inches(0.55)

    def set_cell_background(cell, hex_color):
        tcPr = cell._tc.get_or_add_tcPr()
        shd = OxmlElement('w:shd')
        shd.set(qn('w:val'), 'clear')
        shd.set(qn('w:color'), 'auto')
        shd.set(qn('w:fill'), hex_color)
        tcPr.append(shd)

    # Document Title
    p_title = doc.add_paragraph()
    r_title = p_title.add_run(REPORT_HEADER["title"])
    r_title.bold = True
    r_title.font.size = Pt(17)
    r_title.font.color.rgb = RGBColor(0x1A, 0x2B, 0x4C)
    p_title.paragraph_format.space_after = Pt(2)

    p_sub = doc.add_paragraph()
    r_sub = p_sub.add_run("Undergraduate Technical Review | B.Tech Computer Science | CortexHR Project")
    r_sub.bold = True
    r_sub.font.size = Pt(10)
    r_sub.font.color.rgb = RGBColor(0x2B, 0x6C, 0xB0)
    p_sub.paragraph_format.space_after = Pt(8)

    # Metadata Box
    meta_tbl = doc.add_table(rows=3, cols=4)
    meta_tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    meta_data = [
        ("Student Name:", REPORT_HEADER["name"], "Track / Phase:", f"{REPORT_HEADER['track']} | {REPORT_HEADER['phase']}"),
        ("Project:", REPORT_HEADER["project"], "Git Branch:", f"{REPORT_HEADER['branch']} (Tag: pending)"),
        ("Repository:", REPORT_HEADER["repo"], "Evaluation:", "Labels strictly sealed during training"),
    ]
    for row_idx, row in enumerate(meta_data):
        for col_idx in range(4):
            cell = meta_tbl.cell(row_idx, col_idx)
            cell.text = row[col_idx]
            set_cell_background(cell, "F7FAFC")
            for p in cell.paragraphs:
                p.paragraph_format.space_after = Pt(1)
                p.paragraph_format.space_before = Pt(1)
                for r in p.runs:
                    r.font.size = Pt(8)
                    if col_idx % 2 == 0:
                        r.bold = True
                        r.font.color.rgb = RGBColor(0x1A, 0x2B, 0x4C)

    doc.add_paragraph().paragraph_format.space_after = Pt(4)

    # 1. Summary of Phase 1 Tasks
    h1 = doc.add_heading(level=1)
    r1 = h1.add_run("1. Summary of Phase 1 Tasks")
    r1.font.size = Pt(12)
    r1.font.color.rgb = RGBColor(0x1A, 0x2B, 0x4C)
    h1.paragraph_format.space_before = Pt(6)
    h1.paragraph_format.space_after = Pt(4)

    t_task = doc.add_table(rows=1 + len(TASKS_SUMMARY), cols=4)
    t_task.alignment = WD_TABLE_ALIGNMENT.CENTER
    headers = ["Task", "Status", "What I Worked On", "Verified Evidence"]
    for i, h in enumerate(headers):
        cell = t_task.cell(0, i)
        cell.text = h
        set_cell_background(cell, "1A2B4C")
        for p in cell.paragraphs:
            for r in p.runs:
                r.bold = True
                r.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
                r.font.size = Pt(8)

    for idx, row in enumerate(TASKS_SUMMARY):
        c0 = t_task.cell(idx + 1, 0)
        c0.text = row["task"]
        c1 = t_task.cell(idx + 1, 1)
        c1.text = row["status"]
        c2 = t_task.cell(idx + 1, 2)
        c2.text = row["worked_on"]
        c3 = t_task.cell(idx + 1, 3)
        c3.text = row["evidence"]
        bg = "FFFFFF" if idx % 2 == 0 else "F8FAFC"
        for col_i in range(4):
            c = t_task.cell(idx + 1, col_i)
            set_cell_background(c, bg)
            for p in c.paragraphs:
                p.paragraph_format.space_after = Pt(2)
                p.paragraph_format.space_before = Pt(2)
                for r in p.runs:
                    r.font.size = Pt(8)
                    if col_i == 0:
                        r.bold = True
                    if col_i == 1:
                        r.font.color.rgb = RGBColor(0x0D, 0x68, 0x32)
                        r.bold = True

    # 2. Proof That the Work Runs
    h2 = doc.add_heading(level=1)
    r2 = h2.add_run("2. Proof That the Work Runs & Verified Metrics")
    r2.font.size = Pt(12)
    r2.font.color.rgb = RGBColor(0x1A, 0x2B, 0x4C)
    h2.paragraph_format.space_before = Pt(8)
    h2.paragraph_format.space_after = Pt(4)

    p_runs = doc.add_paragraph()
    p_runs.add_run("All code, unit tests, and validation scripts were executed against real project data. No benchmark figures were copied as my own.")
    p_runs.paragraph_format.space_after = Pt(4)

    # Subsections
    doc.add_heading("A. Data Audit & Feature Pipeline (Tasks 1 & 2)", level=2)
    p_a = doc.add_paragraph()
    p_a.add_run(f"• Dataset Intake: 2,389 raw rows; dropped 6 duplicates ({METRICS_SUMMARY['task1']['duplicates']}) and 4 negative tenure errors ({METRICS_SUMMARY['task1']['invalid_tenure']}).\n")
    p_a.add_run(f"• Target Leakage: Dropped 6 exit-related columns ({METRICS_SUMMARY['task1']['leaks']}) preventing false 100% accuracy.\n")
    p_a.add_run(f"• Pipeline Output: 2,379 clean records (349 leavers, 14.67% base prevalence). Engineered rolling 3-month attendance aggregates joined strictly inside CV folds.")
    p_a.paragraph_format.space_after = Pt(4)

    doc.add_heading("B. Baseline Classification Results (Task 3)", level=2)
    t3 = doc.add_table(rows=1 + len(METRICS_SUMMARY["task3"]), cols=7)
    t3.alignment = WD_TABLE_ALIGNMENT.CENTER
    t3_h = ["Model", "ROC-AUC (Ours)", "PR-AUC", "Precision", "Recall", "Accuracy", "Benchmark"]
    for i, h in enumerate(t3_h):
        cell = t3.cell(0, i)
        cell.text = h
        set_cell_background(cell, "2B6CB0")
        for p in cell.paragraphs:
            for r in p.runs:
                r.bold = True
                r.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
                r.font.size = Pt(7.5)

    for idx, row in enumerate(METRICS_SUMMARY["task3"]):
        for col_i in range(7):
            cell = t3.cell(idx + 1, col_i)
            cell.text = str(row[col_i])
            bg = "FFFFFF" if idx % 2 == 0 else "F8FAFC"
            set_cell_background(cell, bg)
            for p in cell.paragraphs:
                for r in p.runs:
                    r.font.size = Pt(7.5)
                    if col_i == 0:
                        r.bold = True

    doc.add_heading("C. Calibration & Operational Risk Bands (Task 4)", level=2)
    p_c = doc.add_paragraph()
    p_c.add_run(f"• Probability Calibration: Sigmoid scaling reduced Brier score from {METRICS_SUMMARY['task4']['brier_uncalibrated']} to {METRICS_SUMMARY['task4']['brier_calibrated']}, comfortably beating naive prevalence baseline (0.125).\n")
    p_c.add_run(f"• Top-Decile Lift: Top 10% scored employees (238 people) exhibited 2.26x lift (33.2% leaver concentration vs 14.7% base rate), capturing 22.6% of all leavers.\n")
    p_c.add_run(f"• Test Execution: All 29 unit tests in test_evaluation.py passed successfully.")
    p_c.paragraph_format.space_after = Pt(4)

    doc.add_heading("D. Unsupervised Attendance Anomaly Detection (Task 5)", level=2)
    t5 = doc.add_table(rows=1 + len(METRICS_SUMMARY["task5"]), cols=8)
    t5.alignment = WD_TABLE_ALIGNMENT.CENTER
    t5_h = ["Approach", "Precision", "Recall", "F1", "TP", "FP", "FN", "Flagged"]
    for i, h in enumerate(t5_h):
        cell = t5.cell(0, i)
        cell.text = h
        set_cell_background(cell, "1A2B4C")
        for p in cell.paragraphs:
            for r in p.runs:
                r.bold = True
                r.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
                r.font.size = Pt(7.5)

    for idx, row in enumerate(METRICS_SUMMARY["task5"]):
        for col_i in range(8):
            cell = t5.cell(idx + 1, col_i)
            cell.text = str(row[col_i])
            bg = "EFF6FF" if "Combined" in row[0] else ("FFFFFF" if idx % 2 == 0 else "F8FAFC")
            set_cell_background(cell, bg)
            for p in cell.paragraphs:
                for r in p.runs:
                    r.font.size = Pt(7.5)
                    if col_i == 0:
                        r.bold = True

    p_d = doc.add_paragraph()
    p_d.add_run("• Evaluation Protocol: Trained strictly without labels (3% contamination). Combined model + business rules achieved 0.488 Precision and 0.827 Recall vs 0.403/0.593 published baseline.\n")
    p_d.add_run("• Missed Types: 100% recall on Absence Streaks, Impossible Hours, Sudden Drops, and Zero Attendance. 26 missed records were Late Arrival Spikes (single-dimension deviations blending in among 8 normal features).\n")
    p_d.add_run("• Unit Tests: All 14 test cases in tests/test_attendance_anomaly_detector.py passed cleanly.")
    p_d.paragraph_format.space_after = Pt(6)

    # 3. What I Learned
    h3 = doc.add_heading(level=1)
    r3 = h3.add_run("3. What I Learned (Student Reflection)")
    r3.font.size = Pt(12)
    r3.font.color.rgb = RGBColor(0x1A, 0x2B, 0x4C)
    p_learn = doc.add_paragraph()
    r_learn = p_learn.add_run(f'"{LEARNING_REFLECTION}"')
    r_learn.italic = True
    r_learn.font.color.rgb = RGBColor(0x1A, 0x2B, 0x4C)
    p_learn.paragraph_format.space_after = Pt(6)

    # 4. Problems I Faced and How I Solved Them
    h4 = doc.add_heading(level=1)
    r4 = h4.add_run("4. Real Engineering Challenges & Solutions")
    r4.font.size = Pt(12)
    r4.font.color.rgb = RGBColor(0x1A, 0x2B, 0x4C)
    for prob in PROBLEMS_AND_SOLUTIONS:
        doc.add_heading(prob["issue"], level=2)
        p_p = doc.add_paragraph()
        p_p.add_run("Problem: ").bold = True
        p_p.add_run(prob["description"] + "\n")
        p_p.add_run("Investigation & Resolution: ").bold = True
        p_p.add_run(prob["solution"])
        p_p.paragraph_format.space_after = Pt(4)

    # 5. Roadmap for Phase 2
    h5 = doc.add_heading(level=1)
    r5 = h5.add_run("5. Roadmap for Phase 2 (Payroll Anomaly Detection)")
    r5.font.size = Pt(12)
    r5.font.color.rgb = RGBColor(0x1A, 0x2B, 0x4C)
    p_plan_intro = doc.add_paragraph()
    p_plan_intro.add_run("In Phase 2, anomaly detection extends to payroll, where mistakes carry immediate financial consequences. Planned next steps:")
    for plan in PHASE_2_PLAN:
        p_step = doc.add_paragraph()
        p_step.add_run(f"• {plan}")
        p_step.paragraph_format.space_after = Pt(2)

    doc.save(output_path)
    print(f"[OK] DOCX report successfully generated: {output_path}")


if __name__ == "__main__":
    reports_dir = Path(__file__).resolve().parent.parent / "reports"
    reports_dir.mkdir(parents=True, exist_ok=True)

    pdf_out = str(reports_dir / "CortexHR_Phase_1_Review_Report.pdf")
    docx_out = str(reports_dir / "CortexHR_Phase_1_Review_Report.docx")

    generate_pdf(pdf_out)
    generate_docx(docx_out)
